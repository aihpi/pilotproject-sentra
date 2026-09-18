"""Collection templates: the on-ramp for the people writing test cases.

A round is only as good as its cases, and the cases come from WD and Hotline
staff who will not be editing YAML. So: an Excel sheet for collecting many at
once, a Word form for collecting one at a time, and a reader that turns a
filled sheet back into the case file `cli import` already understands.

`FIELDS` is the single definition all three share, and `test_fields_match_the_importer`
asserts it against `CaseEntry`. That test is the point of the module being
shaped this way: a spreadsheet whose columns have quietly drifted from the
importer is worse than no spreadsheet, because the drift is discovered by
somebody who has already filled in two hundred rows.

Two things deliberately not collected:

  Test-ID   allocated by the backend, never reused, "auch bei zurückgezogenen
            Testfällen nicht". A column for it is an invitation to write one by
            hand, which is how a number gets used twice.
  Status    a case is approved by review, not by the person who wrote it. The
            sheet writes `entwurf` and somebody else decides.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from sentra_eval.categories import KATEGORIE_NAMEN, Kategorie

JA = "ja"
NEIN = "nein"


class VorlageError(RuntimeError):
    """A template could not be built or read. Names the row and the column."""


@dataclass(frozen=True)
class Feld:
    """One thing we ask for, and why."""

    name: str
    label: str
    hilfe: str
    pflicht: bool = False
    auswahl: tuple[str, ...] = ()
    breite: int = 40
    mehrzeilig: bool = False


FIELDS: tuple[Feld, ...] = (
    Feld(
        name="kategorie",
        label="Kategorie",
        hilfe=(
            "Grobe Einordnung der Frage. Bestimmt die Test-ID. "
            + "; ".join(f"{k} = {n}" for k, n in KATEGORIE_NAMEN.items())
            + "."
        ),
        pflicht=True,
        auswahl=tuple(k.value for k in Kategorie),
        breite=12,
    ),
    Feld(
        name="ausgangsfrage",
        label="Ausgangsfrage",
        hilfe=(
            "Die Frage so, wie sie tatsächlich gestellt wurde — nicht "
            "nachträglich geglättet. Eine umformulierte Frage prüft etwas "
            "anderes als die echte."
        ),
        pflicht=True,
        breite=55,
        mehrzeilig=True,
    ),
    Feld(
        name="abteilung",
        label="Abteilung / Fachbereich",
        hilfe="Woher die Frage kam, z. B. Hotline oder WD 3.",
        breite=22,
    ),
    Feld(
        name="erwartete_antwort",
        label="Erwartete Antwort",
        hilfe=(
            "Was eine fachlich richtige Antwort enthalten muss. "
            "UNBEDINGT ausfüllen, BEVOR SENTRA die Frage gestellt bekommt: "
            "eine Erwartung, die nach der Ausgabe geschrieben wird, misst "
            "nichts. Stichpunkte genügen. "
            "Bei einem Grenzfall gehört hierhin, dass keine Antwort erwartet "
            "wird und warum."
        ),
        pflicht=True,
        breite=55,
        mehrzeilig=True,
    ),
    Feld(
        name="referenz_korrekt",
        label="Korrekte Quelle",
        hilfe=(
            "Die Fundstelle, auf der die richtige Antwort beruht, so wie ein "
            "Mensch sie liest — z. B. „GOBT § 35“. Pflicht, außer bei einem "
            "Grenzfall: zu einer Frage, die der Bestand nicht abdeckt, gibt es "
            "keine korrekte Quelle."
        ),
        breite=35,
    ),
    Feld(
        name="referenz_korrekt_az",
        label="Aktenzeichen der korrekten Quelle",
        hilfe=(
            "Dasselbe Dokument als Aktenzeichen, genau in der Schreibweise des "
            "Bestands: „WD 3 - 3000 - 029/23“. Nur hierauf prüft die "
            "Quellenprüfung 4.3b automatisch. Leer lassen, wenn es zu der Frage "
            "kein Dokument im Bestand gibt — die Prüfung meldet dann "
            "„nicht prüfbar“ statt eines falschen Befunds."
        ),
        breite=30,
    ),
    Feld(
        name="referenz_falsch",
        label="Veraltete / ähnliche Quelle",
        hilfe=(
            "Optional, aber wertvoll: eine thematisch naheliegende, aber "
            "veraltete oder falsche Fundstelle. Erst dadurch unterscheidet die "
            "Prüfung „hat eine Quelle genannt“ von „hat die richtige von zwei "
            "plausiblen genannt“."
        ),
        breite=35,
    ),
    Feld(
        name="referenz_falsch_az",
        label="Aktenzeichen der veralteten Quelle",
        hilfe=(
            "Dieselbe veraltete Quelle als Aktenzeichen, in der Schreibweise "
            "des Bestands. Nur hierauf prüft 4.3b, ob SENTRA die veraltete "
            "statt der gültigen Fassung herangezogen hat — der Befund, für den "
            "die beiden Quellenfelder überhaupt da sind."
        ),
        breite=30,
    ),
    Feld(
        name="grund_fuer_aufnahme",
        label="Grund für die Aufnahme",
        hilfe=(
            "Warum dieser Fall in den Datensatz gehört — häufige Frage, "
            "bekannte Schwachstelle, Beschwerde, Grenzfall. Steht später im "
            "Dokumentationsbogen und erklärt dem Prüfenden, worauf zu achten "
            "ist."
        ),
        breite=45,
        mehrzeilig=True,
    ),
    Feld(
        name="grenzfall",
        label="Grenzfall?",
        hilfe=(
            "„ja“, wenn der Bestand die Frage bewusst nicht abdeckt und SENTRA "
            "die Antwort verweigern soll. Grenzfälle werden immer manuell "
            "geprüft und nie automatisch herausgefiltert — sie sind die "
            "riskanteste Kategorie: ein System, das etwas erfindet, ist "
            "schlimmer als eines, das nichts findet."
        ),
        auswahl=(NEIN, JA),
        breite=12,
    ),
)

FIELD_NAMES: tuple[str, ...] = tuple(f.name for f in FIELDS)

TITEL = "SENTRA — Erfassung von Testfällen"

EINLEITUNG = (
    "Mit diesem Formular werden die Testfälle gesammelt, gegen die SENTRA "
    "regelmäßig geprüft wird. Jede Zeile ist ein Testfall: eine echte Frage, "
    "die erwartete Antwort und die Quelle, auf der sie beruht."
)

DIE_EINE_REGEL = (
    "Die wichtigste Regel: Die erwartete Antwort muss feststehen, bevor SENTRA "
    "die Frage gestellt bekommt. Wird sie nachträglich an die tatsächliche "
    "Ausgabe angepasst, misst der Testfall nichts mehr."
)

NICHT_AUSFUELLEN = (
    "Test-ID und Status werden nicht hier vergeben. Die Test-ID vergibt das "
    "System und benutzt sie nie ein zweites Mal; über die Freigabe entscheidet "
    "die fachliche Durchsicht, nicht die erfassende Person."
)


# ── The Excel sheet ─────────────────────────────────────────────────


def _require_openpyxl() -> Any:
    """openpyxl, or a sentence saying how to get it.

    The bare ImportError names a module, which is no help to the operator
    importing a sheet somebody sent them — they did not choose to depend on
    openpyxl and have no reason to know the extra exists.
    """
    try:
        import openpyxl
    except ModuleNotFoundError as exc:  # pragma: no cover - depends on the install
        raise VorlageError(
            "Für Excel-Vorlagen fehlt openpyxl. Installieren mit: uv sync --extra vorlagen"
        ) from exc
    return openpyxl


def _require_docx() -> Any:
    try:
        import docx
    except ModuleNotFoundError as exc:  # pragma: no cover - depends on the install
        raise VorlageError(
            "Für die Word-Vorlage fehlt python-docx. Installieren mit: uv sync --extra vorlagen"
        ) from exc
    return docx


def build_workbook() -> Any:
    """A sheet for collecting many cases at once.

    Two tabs. The first is what gets filled in, the second explains every
    column — because a header row cannot hold "write this before SENTRA sees
    the question and not after", and that is the sentence the whole exercise
    depends on.

    Categories and Grenzfall are dropdowns rather than free text. Not tidiness:
    `kategorie` decides the Test-ID, and a typo there is a case that cannot be
    imported at all.
    """
    _require_openpyxl()
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Testfälle"

    kopf_fill = PatternFill("solid", fgColor="1F3864")
    kopf_font = Font(color="FFFFFF", bold=True, size=10)
    pflicht_fill = PatternFill("solid", fgColor="C00000")
    rand = Border(bottom=Side(style="thin", color="BFBFBF"))

    for index, feld in enumerate(FIELDS, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.value = feld.label + (" *" if feld.pflicht else "")
        cell.fill = pflicht_fill if feld.pflicht else kopf_fill
        cell.font = kopf_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        # The help text as a cell comment, so it is reachable where it is
        # needed rather than only on the other tab.
        cell.comment = _comment(feld)
        sheet.column_dimensions[get_column_letter(index)].width = feld.breite

    sheet.row_dimensions[1].height = 34
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(FIELDS))}1"

    # Room to fill in, pre-formatted. Without this, wrapped text in a long
    # expected answer collapses to one line and people write less than they
    # would have.
    for row in range(2, 302):
        sheet.row_dimensions[row].height = 30
        for index, feld in enumerate(FIELDS, start=1):
            cell = sheet.cell(row=row, column=index)
            cell.alignment = Alignment(vertical="top", wrap_text=feld.mehrzeilig)
            cell.border = rand

    for index, feld in enumerate(FIELDS, start=1):
        if not feld.auswahl:
            continue
        letter = get_column_letter(index)
        validation = DataValidation(
            type="list",
            formula1='"' + ",".join(feld.auswahl) + '"',
            allow_blank=True,
            showDropDown=False,
        )
        validation.error = "Bitte einen Wert aus der Liste wählen."
        validation.errorTitle = feld.label
        validation.prompt = feld.hilfe
        validation.promptTitle = feld.label
        sheet.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}301")

    _add_guide(workbook)
    return workbook


def _comment(feld: Feld) -> Any:
    from openpyxl.comments import Comment

    comment = Comment(feld.hilfe, "SENTRA")
    comment.width = 320
    comment.height = 160
    return comment


def _add_guide(workbook: Any) -> None:
    """The tab that explains the columns, and the rule behind all of them."""
    from openpyxl.styles import Alignment, Font

    guide = workbook.create_sheet("Ausfüllhinweise")
    guide.column_dimensions["A"].width = 34
    guide.column_dimensions["B"].width = 96

    def schreibe(row: int, links: str, rechts: str, *, fett: bool = False) -> None:
        guide.cell(row=row, column=1, value=links).font = Font(bold=True, size=10)
        cell = guide.cell(row=row, column=2, value=rechts)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(bold=fett, size=10)
        guide.row_dimensions[row].height = max(30, 15 * (len(rechts) // 90 + 1))

    guide["A1"] = TITEL
    guide["A1"].font = Font(bold=True, size=14)

    schreibe(3, "Worum es geht", EINLEITUNG)
    schreibe(4, "Die wichtigste Regel", DIE_EINE_REGEL, fett=True)
    schreibe(5, "Nicht ausfüllen", NICHT_AUSFUELLEN)
    schreibe(6, "Pflichtfelder", "Rot hinterlegt und mit * markiert.")

    guide["A8"] = "Die Felder im Einzelnen"
    guide["A8"].font = Font(bold=True, size=12)

    for offset, feld in enumerate(FIELDS):
        schreibe(9 + offset, feld.label + (" *" if feld.pflicht else ""), feld.hilfe)


# ── The Word form ───────────────────────────────────────────────────


def build_document() -> Any:
    """A form for collecting one case at a time.

    The counterpart to the sheet rather than a copy of it: somebody writing up
    a single question they were just asked will not open a spreadsheet, and a
    form with the guidance beside each field gets better expected answers than
    a row with a comment on the header.

    The input areas are real Word content controls, so the form can be tabbed
    through and the explanatory text is not something to accidentally type
    over. python-docx has no API for them, so they are built as XML — a
    plain-text `w:sdt` per field, which is the simplest control Word offers and
    the one that round-trips everywhere.
    """
    _require_docx()
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    document = Document()

    title = document.add_heading(TITEL, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    intro = document.add_paragraph(EINLEITUNG)
    intro.runs[0].font.size = Pt(10)

    regel = document.add_paragraph()
    run = regel.add_run(DIE_EINE_REGEL)
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    hinweis = document.add_paragraph(NICHT_AUSFUELLEN)
    hinweis.runs[0].font.size = Pt(9)
    hinweis.runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    document.add_paragraph()

    for feld in FIELDS:
        label = document.add_paragraph()
        run = label.add_run(feld.label + (" *" if feld.pflicht else ""))
        run.bold = True
        run.font.size = Pt(10)

        hilfe = document.add_paragraph()
        hilfe_run = hilfe.add_run(feld.hilfe)
        hilfe_run.italic = True
        hilfe_run.font.size = Pt(8.5)
        hilfe_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

        _add_control(document, feld)
        document.add_paragraph()

    footer = document.add_paragraph()
    footer_run = footer.add_run(
        "Ausgefülltes Formular an die für die Auswertung zuständige Person. "
        "Mehrere Fälle auf einmal: bitte die Excel-Vorlage benutzen."
    )
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    return document


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _add_control(document: Any, feld: Feld) -> None:
    """One content control, shaded, as a paragraph of its own.

    A dropdown where the field has a fixed set of answers and a plain text box
    otherwise. The dropdown matters for the same reason it does in the sheet:
    `kategorie` decides the Test-ID, and a typo is a case that cannot be
    imported.
    """
    from docx.oxml.ns import qn
    from docx.oxml.parser import OxmlElement

    paragraph = document.add_paragraph()
    # Shaded, so the space to write in is visible before anything is typed.
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), "F2F2F2")
    paragraph.paragraph_format.element.get_or_add_pPr().append(shading)

    sdt = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")

    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), feld.label)
    properties.append(alias)

    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), feld.name)
    properties.append(tag)

    placeholder = OxmlElement("w:showingPlcHdr")
    properties.append(placeholder)

    if feld.auswahl:
        dropdown = OxmlElement("w:dropDownList")
        for option in feld.auswahl:
            item = OxmlElement("w:listItem")
            item.set(qn("w:displayText"), option)
            item.set(qn("w:value"), option)
            dropdown.append(item)
        properties.append(dropdown)
    else:
        text = OxmlElement("w:text")
        if feld.mehrzeilig:
            text.set(qn("w:multiLine"), "1")
        properties.append(text)

    sdt.append(properties)

    content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    run_text = OxmlElement("w:t")
    run_text.text = feld.auswahl[0] if feld.auswahl else "Klicken und eintragen"
    run.append(run_text)
    content.append(run)
    sdt.append(content)

    # The sdt replaces the paragraph's empty run, wrapping the content inline.
    paragraph.paragraph_format.element.append(sdt)


# ── Reading a filled sheet back ─────────────────────────────────────


def read_workbook(path: Path) -> list[dict[str, Any]]:
    """A filled sheet as case entries, ready for `yaml_io.parse`.

    Matches columns by their header rather than by position, so a sheet that
    somebody reordered or added a working column to still reads. A header we do
    not recognise is ignored; a missing required one is an error naming it,
    because silently importing without `ausgangsfrage` would create cases that
    ask nothing.

    Rows are numbered as Excel numbers them in the error messages. Telling
    somebody "row 3" when they are looking at row 4 is worse than not telling
    them.
    """
    _require_openpyxl()
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    sheet = workbook["Testfälle"] if "Testfälle" in workbook.sheetnames else workbook.worksheets[0]

    header = [_norm(cell.value) for cell in sheet[1]]
    by_label = {_norm(f.label): f for f in FIELDS}
    columns: dict[int, Feld] = {}
    for index, title in enumerate(header):
        feld = by_label.get(title)
        if feld is not None:
            columns[index] = feld

    fehlend = [f.label for f in FIELDS if f.pflicht and f not in columns.values()]
    if fehlend:
        raise VorlageError(
            f"{path}: die Pflichtspalte(n) {', '.join(fehlend)} fehlen. "
            "Wurde die Kopfzeile verändert?"
        )

    entries: list[dict[str, Any]] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        werte = {
            feld.name: _clean(row[index].value)
            for index, feld in columns.items()
            if index < len(row)
        }
        # An untouched row is not an error. A sheet handed out with 300 blank
        # rows and eleven filled in is the normal case.
        if not any(werte.values()):
            continue

        entry: dict[str, Any] = {name: werte.get(name, "") for name in FIELD_NAMES}
        entry["kategorie"] = entry["kategorie"].upper()
        entry["grenzfall"] = _bool(entry["grenzfall"], row_number)
        # Never from the sheet. The backend allocates the one and review
        # decides the other.
        entry["status"] = "entwurf"

        _check_required(entry, row_number, path)
        entries.append(entry)

    return entries


def _check_required(entry: dict[str, Any], row_number: int, path: Path) -> None:
    for feld in FIELDS:
        if feld.pflicht and not str(entry.get(feld.name, "")).strip():
            raise VorlageError(f"{path} Zeile {row_number}: {feld.label} fehlt.")

    kategorie = entry["kategorie"]
    erlaubt = {k.value for k in Kategorie}
    if kategorie not in erlaubt:
        raise VorlageError(
            f"{path} Zeile {row_number}: Kategorie „{kategorie}“ ist unbekannt. "
            f"Erlaubt: {', '.join(sorted(erlaubt))}."
        )

    # The rule the case store enforces at approval, checked here so it is
    # found while the sheet is open rather than at import time.
    if not entry["grenzfall"] and not str(entry["referenz_korrekt"]).strip():
        raise VorlageError(
            f"{path} Zeile {row_number}: Ohne „{_label('referenz_korrekt')}“ kann der "
            "Testfall später nicht freigegeben werden. Bei einer Frage, die der Bestand "
            "bewusst nicht abdeckt, bitte Grenzfall auf „ja“ setzen."
        )


def _label(name: str) -> str:
    return next(f.label for f in FIELDS if f.name == name)


def _norm(value: Any) -> str:
    return str(value or "").replace("*", "").strip().casefold()


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _bool(value: str, row_number: int) -> bool:
    text = value.strip().casefold()
    if text in {"", NEIN, "nein.", "no", "false", "0", "n"}:
        return False
    if text in {JA, "ja.", "yes", "true", "1", "x", "j"}:
        return True
    raise VorlageError(f"Zeile {row_number}: „{value}“ ist kein ja/nein für Grenzfall.")


# ── Writing the templates out ───────────────────────────────────────

WORKBOOK_NAME = "SENTRA-Testfaelle-Erfassung.xlsx"
DOCUMENT_NAME = "SENTRA-Testfall-Erfassung.docx"


def write_templates(directory: Path) -> list[Path]:
    """Both templates, regenerated from FIELDS.

    Committed to the repository as well as generated, because the people who
    need them do not have a checkout — but generated, so that adding a field
    is one change here rather than three files edited by hand and two of them
    forgotten.
    """
    directory.mkdir(parents=True, exist_ok=True)
    workbook_path = directory / WORKBOOK_NAME
    document_path = directory / DOCUMENT_NAME

    build_workbook().save(workbook_path)
    build_document().save(document_path)
    return [workbook_path, document_path]
