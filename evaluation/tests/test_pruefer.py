"""Reading the Word table the WD reviewers fill in (#244).

The fixture is built here with made-up content in the reviewers' layout, rather
than taken from a real submission: those stay out of the repository.
"""

from pathlib import Path

import pytest

from sentra_eval import cli, pruefer, yaml_io
from sentra_eval.categories import Kategorie
from sentra_eval.vorlagen import VorlageError

HEADER = (
    "Test-ID",
    "Ausgangsfrage",
    "Referenzquelle(korrekt)",
    "Referenzquelle",
    "Referenzquelle (thematisch ähnlich/veraltet/falsch)",
)

ZWEI_FRAGEN = (
    "XY01",
    ["Allgemein:", "Gibt es eine Steuer auf Mondgestein?", "", "", "Konkret:", "Wie hoch ist sie?"],
    "Die Mondgesteinsteuer beträgt 3 Prozent.",
    "Besteuerung von Mondgestein, WD 4 - 3000 - 101/24, S. 5",
    "WD-4-007-19.pdf",
)


def _document(tmp_path: Path, rows: list[tuple], header: tuple = HEADER) -> Path:
    """The reviewers' document: some instructions, then the one table."""
    pytest.importorskip("docx", reason="needs `uv sync --extra vorlagen`")
    from docx import Document

    document = Document()
    document.add_paragraph("Projekt SENTRA - Testverfahren")
    document.add_paragraph("Bitte die Tabelle ausfüllen.")
    table = document.add_table(rows=1, cols=len(header))
    for cell, title in zip(table.rows[0].cells, header, strict=True):
        cell.text = title
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            paragraphs = value if isinstance(value, list) else [value]
            cell.text = paragraphs[0]
            for paragraph in paragraphs[1:]:
                cell.add_paragraph(paragraph)
    path = tmp_path / "pruefer.docx"
    document.save(path)
    return path


def _read(path: Path) -> list[yaml_io.CaseEntry]:
    return yaml_io.validate(pruefer.read_document(path))


class TestReadingARow:
    def test_each_question_becomes_its_own_case(self, tmp_path):
        entries = _read(_document(tmp_path, [ZWEI_FRAGEN]))

        assert [e.ausgangsfrage for e in entries] == [
            "Gibt es eine Steuer auf Mondgestein?",
            "Wie hoch ist sie?",
        ]

    def test_both_share_the_answer_and_the_sources(self, tmp_path):
        allgemein, konkret = _read(_document(tmp_path, [ZWEI_FRAGEN]))

        for entry in (allgemein, konkret):
            assert entry.erwartete_antwort == "Die Mondgesteinsteuer beträgt 3 Prozent."
            assert entry.referenz_korrekt == ZWEI_FRAGEN[3]
            assert entry.referenz_korrekt_az == "WD 4 - 3000 - 101/24"
            assert entry.referenz_falsch == "WD-4-007-19.pdf"
            assert entry.referenz_falsch_az == "WD 4 - 3000 - 007/19"

    def test_the_category_comes_from_the_papers_fachbereich(self, tmp_path):
        wd5 = ("XY02", "Was ist Strom?", "Fließende Ladung.", "WD 5 - 3000 - 002/26", "")

        fi, _, wi = _read(_document(tmp_path, [ZWEI_FRAGEN, wd5]))

        assert fi.kategorie == Kategorie.FI
        assert wi.kategorie == Kategorie.WI

    def test_the_reviewers_id_is_kept_for_tracing_but_is_not_the_test_id(self, tmp_path):
        allgemein, konkret = _read(_document(tmp_path, [ZWEI_FRAGEN]))

        assert allgemein.test_id is None
        assert allgemein.grund_fuer_aufnahme == "Prüfer-Vorlage (XY01, Allgemein)"
        assert konkret.grund_fuer_aufnahme == "Prüfer-Vorlage (XY01, Konkret)"

    def test_nothing_arrives_approved(self, tmp_path):
        entries = _read(_document(tmp_path, [ZWEI_FRAGEN]))

        assert {e.status for e in entries} == {"entwurf"}
        assert not any(e.grenzfall for e in entries)

    def test_empty_rows_are_skipped(self, tmp_path):
        entries = _read(_document(tmp_path, [ZWEI_FRAGEN, ("", "", "", "", "")]))

        assert len(entries) == 2

    def test_the_wrong_source_column_is_optional(self, tmp_path):
        path = _document(tmp_path, [ZWEI_FRAGEN[:4]], header=HEADER[:4])

        assert {e.referenz_falsch_az for e in _read(path)} == {""}


class TestRefusing:
    def test_a_wrong_source_that_is_the_correct_one_names_the_row(self, tmp_path):
        """The 4.3b check could never pass, and nobody would find out why
        until the round's report."""
        row = (*ZWEI_FRAGEN[:4], "WD-4-101-24")

        with pytest.raises(VorlageError, match=r"Zeile 2 \(XY01\).*dieselbe Arbeit"):
            pruefer.read_document(_document(tmp_path, [row]))

    def test_a_fachbereich_without_a_category_names_the_row(self, tmp_path):
        row = ("XY03", "Frage?", "Antwort.", "WD 3 - 3000 - 001/25", "")

        with pytest.raises(VorlageError, match=r"Zeile 2 \(XY03\).*WD 3 - 3000 - 001/25"):
            pruefer.read_document(_document(tmp_path, [row]))

    def test_a_source_without_an_aktenzeichen_names_the_row(self, tmp_path):
        row = ("XY04", "Frage?", "Antwort.", "Irgendein Gutachten", "")

        with pytest.raises(VorlageError, match=r"Zeile 2 \(XY04\).*kein Aktenzeichen"):
            pruefer.read_document(_document(tmp_path, [row]))

    def test_a_document_without_the_table_says_so(self, tmp_path):
        path = _document(tmp_path, [("a", "b")], header=("Name", "Wert"))

        with pytest.raises(VorlageError, match="Ausgangsfrage"):
            pruefer.read_document(path)


class TestSplittingQuestions:
    def test_a_cell_without_markers_is_one_question(self):
        assert pruefer.split_questions("Was ist Strom?") == [("", "Was ist Strom?")]

    def test_markers_on_one_line(self):
        text = "Allgemein: Gibt es das? Konkret: Wie viel?"

        assert pruefer.split_questions(text) == [
            ("Allgemein", "Gibt es das?"),
            ("Konkret", "Wie viel?"),
        ]

    def test_the_order_is_the_reviewers(self):
        text = "Konkret: Wie viel?\nAllgemein: Gibt es das?"

        assert [art for art, _ in pruefer.split_questions(text)] == ["Konkret", "Allgemein"]

    def test_several_questions_under_one_marker_stay_together(self):
        text = "Konkret: Wann war das?\nWer war es?"

        assert pruefer.split_questions(text) == [("Konkret", "Wann war das? Wer war es?")]


class TestAktenzeichen:
    @pytest.mark.parametrize(
        ("text", "expected"),
        [
            ("Titel, WD 4 - 3000 - 115/22", "WD 4 - 3000 - 115/22"),
            ("WD 5 - 3000 – 025/26, S. 9", "WD 5 - 3000 - 025/26"),
            ("WD-6-054-24.pdf", "WD 6 - 3000 - 054/24"),
            ("EU-6-042-24", "EU 6 - 3000 - 042/24"),
            ("WD-4-013-25\n\nWD-4-028-25 + WD-3-058-25", "WD 4 - 3000 - 013/25"),
            ("WD 4 - 3000 - 011/26, WD 5 - 3000 - 017/26", "WD 4 - 3000 - 011/26"),
            ("", ""),
            ("keins", ""),
        ],
    )
    def test_the_first_one_in_canonical_form(self, text, expected):
        assert pruefer.first_aktenzeichen(text) == expected


def test_cli_import_reads_a_docx(tmp_path):
    entries = cli._read(_document(tmp_path, [ZWEI_FRAGEN]))

    assert len(entries) == 2
